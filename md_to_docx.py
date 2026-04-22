#!/usr/bin/env python3
"""Convert the client_analysis.md to a formatted Word document."""

import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def parse_markdown_to_sections(md_text: str) -> list[dict]:
    """Parse markdown into a list of content blocks."""
    blocks = []
    lines = md_text.splitlines()
    i = 0

    while i < len(lines):
        line = lines[i]

        # H1
        if line.startswith("# ") and not line.startswith("## "):
            blocks.append({"type": "h1", "text": line[2:].strip()})
            i += 1
            continue

        # H2
        if line.startswith("## "):
            blocks.append({"type": "h2", "text": line[3:].strip()})
            i += 1
            continue

        # Horizontal rule
        if line.strip() == "---":
            blocks.append({"type": "hr"})
            i += 1
            continue

        # Table
        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            blocks.append({"type": "table", "lines": table_lines})
            continue

        # Bullet list
        if line.startswith("- "):
            items = []
            while i < len(lines) and lines[i].startswith("- "):
                items.append(lines[i][2:].strip())
                i += 1
            blocks.append({"type": "list", "items": items})
            continue

        # Bold paragraph like **Contacts:**
        if line.startswith("**") and line.endswith("**"):
            blocks.append({"type": "bold", "text": line.strip("* ")})
            i += 1
            continue

        # Italic line
        if line.startswith("*") and line.endswith("*") and not line.startswith("**"):
            blocks.append({"type": "italic", "text": line.strip("* ")})
            i += 1
            continue

        # Regular paragraph (non-empty)
        if line.strip():
            blocks.append({"type": "paragraph", "text": line.strip()})

        i += 1

    return blocks


def add_formatted_run(paragraph, text: str):
    """Add text with inline bold/link markdown rendered."""
    # Process **bold** and [text](url) patterns
    pattern = r"(\*\*(.+?)\*\*|\[(.+?)\]\((.+?)\))"
    last_end = 0
    for m in re.finditer(pattern, text):
        # Add text before this match
        if m.start() > last_end:
            paragraph.add_run(text[last_end:m.start()])

        if m.group(2):  # Bold
            run = paragraph.add_run(m.group(2))
            run.bold = True
        elif m.group(3):  # Link - just show the text
            run = paragraph.add_run(m.group(3))
            run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)

        last_end = m.end()

    # Remaining text
    if last_end < len(text):
        paragraph.add_run(text[last_end:])


def parse_table(table_lines: list[str]) -> tuple[list[str], list[list[str]]]:
    """Parse markdown table into headers and rows."""
    headers = []
    rows = []
    for idx, line in enumerate(table_lines):
        cells = [c.strip() for c in line.strip("|").split("|")]
        if idx == 0:
            headers = cells
        elif idx == 1:
            # Separator row (|---|---|...)
            continue
        else:
            rows.append(cells)
    return headers, rows


def build_docx(md_path: Path, output_path: Path):
    md_text = md_path.read_text(encoding="utf-8")
    blocks = parse_markdown_to_sections(md_text)

    doc = Document()

    # Set default font
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(10)

    # Style headings
    for level, size in [(1, 20), (2, 14)]:
        h_style = doc.styles[f"Heading {level}"]
        h_style.font.name = "Calibri"
        h_style.font.size = Pt(size)
        h_style.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    for block in blocks:
        btype = block["type"]

        if btype == "h1":
            doc.add_heading(block["text"], level=1)

        elif btype == "h2":
            doc.add_heading(block["text"], level=2)

        elif btype == "hr":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("─" * 80)
            run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
            run.font.size = Pt(6)

        elif btype == "italic":
            p = doc.add_paragraph()
            run = p.add_run(block["text"])
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

        elif btype == "bold":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            run = p.add_run(block["text"])
            run.bold = True

        elif btype == "paragraph":
            p = doc.add_paragraph()
            add_formatted_run(p, block["text"])

        elif btype == "list":
            for item in block["items"]:
                p = doc.add_paragraph(style="List Bullet")
                add_formatted_run(p, item)

        elif btype == "table":
            headers, rows = parse_table(block["lines"])
            if not headers:
                continue

            table = doc.add_table(rows=1 + len(rows), cols=len(headers))
            table.style = "Light Grid Accent 1"
            table.alignment = WD_TABLE_ALIGNMENT.LEFT

            # Headers
            for j, h in enumerate(headers):
                cell = table.rows[0].cells[j]
                cell.text = ""
                p = cell.paragraphs[0]
                run = p.add_run(h)
                run.bold = True
                run.font.size = Pt(9)

            # Data rows
            for i, row in enumerate(rows):
                for j, val in enumerate(row):
                    if j < len(headers):
                        cell = table.rows[i + 1].cells[j]
                        cell.text = ""
                        p = cell.paragraphs[0]
                        run = p.add_run(val)
                        run.font.size = Pt(9)

            doc.add_paragraph()  # spacing after table

    doc.save(str(output_path))
    print(f"Saved: {output_path}")
    print(f"Size: {output_path.stat().st_size / 1024:.0f} KB")


if __name__ == "__main__":
    script_dir = Path(__file__).parent
    md_path = script_dir / "client_analysis.md"
    output_path = script_dir / "client_analysis.docx"
    build_docx(md_path, output_path)
